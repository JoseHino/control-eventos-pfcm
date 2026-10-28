# -*- coding: utf-8 -*-
"""Consolida los libros Control_Eventos de PFCM en un unico JSON.

Los libros se reinician cada dos anios (el de 2026 arranca la BD en 2025), asi que
la historia completa solo existe fusionando todos los que haya en la carpeta.
Regla: gana el libro mas reciente, y se anota en CONFLICTO que campos diferian.
"""
import collections, datetime, glob, io, json, os, re, sys, warnings
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
warnings.filterwarnings("ignore")

# "Ultimo" es donde PFCM deja los libros nuevos desde octubre de 2026.
ORIGENES = [r"C:\AMMA\Clientes\PFCM\IA", r"C:\AMMA\Clientes\PFCM\Ultimo"]
AQUI = os.path.dirname(os.path.abspath(__file__))
DESTINO = os.path.join(AQUI, "eventos.json")

COLS = ['CODIGO', 'TITULO', 'FECHA_INICIO', 'FECHA_FIN', 'MES', 'TIPO', 'MODALIDAD',
        'TIPO1', 'CARACTER', 'ORGANIZADOR', 'DELEGACION', 'FACTURABLE', 'FACT_ESPACIO',
        'FACT_SERVICIOS', 'FACT_TOTAL', 'ASISTENTES', 'ANULADO']


def orden(ruta):
    """Los nombres llevan un sello tipo (260105). Si falta, se usa la fecha del fichero.

    OneDrive renombra los ficheros por su cuenta (anade " (1)"), asi que nunca se
    buscan por nombre exacto: se localizan por patron y se ordenan por ese sello.
    """
    m = re.search(r"\((\d{6})\)", os.path.basename(ruta))
    sello = m.group(1) if m else datetime.date.fromtimestamp(os.path.getmtime(ruta)).strftime("%y%m%d")
    return (sello, os.path.getmtime(ruta))


def texto(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        s = " ".join(v.split()).strip()
        return s if s and s != '---' else None
    return v


def numero(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


libros = sorted((l for o in ORIGENES for l in glob.glob(os.path.join(o, "*.xlsm"))), key=orden)
if not libros:
    # Una lectura fallida no debe dejar el destino vacio: se aborta sin escribir.
    sys.exit("ERROR: ningun .xlsm en " + ", ".join(ORIGENES) + ". No se toca " + DESTINO)
print("libros (de mas antiguo a mas reciente):")
for l in libros:
    print("  -", os.path.basename(l))

recs, orden_claves = {}, []
for fi, ruta in enumerate(libros):
    hoja = openpyxl.load_workbook(ruta, data_only=True)["BD"]
    for r in range(3, hoja.max_row + 1):
        d = dict(zip(COLS, [hoja.cell(r, c).value for c in range(1, 18)]))
        if not d['CODIGO'] and not d['TITULO']:
            continue
        # Filas de rotulo al final del formato ("aqui se acaba el formato"): sin codigo ni fecha.
        if not d['CODIGO'] and not d['FECHA_INICIO']:
            continue
        rec = {k: texto(v) for k, v in d.items()}
        for k in ('FACT_ESPACIO', 'FACT_SERVICIOS', 'FACT_TOTAL', 'ASISTENTES'):
            rec[k] = numero(d[k])
        rec['CODIGO'] = str(d['CODIGO']).strip() if d['CODIGO'] else None
        rec['ORIGEN'] = os.path.basename(ruta)
        clave = rec['CODIGO'] or "__sin_codigo_%d_%d" % (fi, r)

        if clave not in recs:
            recs[clave] = rec
            orden_claves.append(clave)
        elif recs[clave]['ORIGEN'] == rec['ORIGEN']:
            # Mismo codigo repetido dentro del mismo libro: son eventos distintos con
            # un codigo mal tecleado. Se conservan para que el aviso los senale.
            # (Antes se miraba "no es el ultimo libro", y con tres libros duplicaba
            # los solapes entre los dos primeros.)
            recs[clave + '#' + str(r)] = rec
            orden_claves.append(clave + '#' + str(r))
        else:
            previo = recs[clave]
            difieren = [c for c in ('TITULO', 'FECHA_INICIO', 'FECHA_FIN', 'TIPO',
                                    'CARACTER', 'FACTURABLE', 'ASISTENTES', 'ANULADO',
                                    'DELEGACION') if previo.get(c) != rec.get(c)]
            rec['CONFLICTO'] = difieren or None
            recs[clave] = rec

data = [recs[k] for k in orden_claves]
data.sort(key=lambda d: (d.get('FECHA_INICIO') or '9999', d.get('CODIGO') or ''))

print("\nregistros consolidados:", len(data))
print("con conflicto entre libros:", sum(1 for d in data if d.get('CONFLICTO')))
print("por anio:", dict(sorted(collections.Counter(
    (d['FECHA_INICIO'] or '????')[:4] for d in data).items())))
print("sin fecha de inicio:", sum(1 for d in data if not d['FECHA_INICIO']))

io.open(DESTINO, 'w', encoding='utf-8', newline='\n').write(
    json.dumps(data, ensure_ascii=False, separators=(',', ':')))
print("escrito:", DESTINO, os.path.getsize(DESTINO), "bytes")
